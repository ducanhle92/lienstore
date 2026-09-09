"""Import / update the catalogue in data/seed.json from an Excel workbook produced by export_products_xlsx.py
(or any sheet with the same column names).

    python scripts/xlsx/import_products_xlsx.py lienstore-san-pham.xlsx [--sheet "Sản phẩm"] [--dry-run]

Behaviour
  • Rows are matched to existing products by ID, then by slug; unmatched rows become new products (ids continue from max).
  • Categories are matched by name (case/space-insensitive) or slug; sheet "Danh mục" (if present) can add new categories.
  • http(s) image URLs are downloaded to public/sites/lienstore/shared/products/import/<slug>[-n].jpg with a 300x300 thumb;
    local paths (/sites/...) are kept as-is. Already-downloaded files are reused.
  • Empty cells never erase existing values, except "Giá gốc"/"Giá vốn"/"Tồn kho" when the cell contains "-" (explicit clear).
  • Rows without a sale price are forced to draft. Nothing is ever deleted.
  • Finally bumps seed.meta.seededAt so running instances pick the change up (LIEN_SEED_SYNC=update recommended).
"""
import argparse, datetime, io, json, os, re, sys, time, unicodedata
import openpyxl, requests
from PIL import Image, ImageOps

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
SEED = os.path.join(ROOT, "data", "seed.json")
IMG_DIR = os.path.join(ROOT, "public", "sites", "lienstore", "shared", "products", "import")
WEB_DIR = "/sites/lienstore/shared/products/import"

ap = argparse.ArgumentParser()
ap.add_argument("xlsx"); ap.add_argument("--sheet", default=None); ap.add_argument("--dry-run", action="store_true")
ap.add_argument("--diff", action="store_true", help="print changed fields per product")
args = ap.parse_args()

def norm(s):  # accent/case/space-insensitive key
    s = unicodedata.normalize("NFD", str(s or "").lower()).replace("đ", "d")
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    return re.sub(r"\s+", " ", s).strip()

def slugify(s):
    s = norm(s); s = re.sub(r"[^a-z0-9]+", "-", s).strip("-")
    return s[:80] or "san-pham"

def to_int(v):
    if v is None: return None
    s = str(v).strip()
    if s in ("", "-"): return None
    digits = re.sub(r"[^\d]", "", s)
    return int(digits) if digits else None

def col(headers, *names):
    """Find a column index by (normalised) header prefix."""
    for n in names:
        for i, h in enumerate(headers):
            if norm(h).startswith(norm(n)): return i
    return None

seed = json.load(open(SEED, encoding="utf-8"))
products = seed["products"]; categories = seed["categories"]
by_id = {p["id"]: p for p in products}; by_slug = {p["slug"]: p for p in products}
cat_by_key = {}
for c in categories:
    cat_by_key[norm(c["name"])] = c["slug"]; cat_by_key[c["slug"]] = c["slug"]

if args.xlsx.lower().endswith(".csv"):
    # CSV fallback (UTF-8, ";" separated, quoted multi-line cells) → wrap it in an in-memory workbook
    import csv
    wb = openpyxl.Workbook(); ws0 = wb.active; ws0.title = "Sản phẩm"
    with open(args.xlsx, encoding="utf-8-sig", newline="") as f:
        sample = f.read(4096); f.seek(0)
        delim = ";" if sample.count(";") >= sample.count(",") else ","
        for row in csv.reader(f, delimiter=delim):
            ws0.append(row)
else:
    wb = openpyxl.load_workbook(args.xlsx, data_only=True)

# --- categories sheet (optional)
if "Danh mục" in wb.sheetnames:
    ws = wb["Danh mục"]; rows = list(ws.iter_rows(values_only=True))
    hi = next((i for i, r in enumerate(rows) if any(norm(x) in ("danh muc", "ten danh muc", "ten danh muc *") for x in r if x)), None)
    if hi is not None:
        h = [str(x or "") for x in rows[hi]]
        ci = {"name": col(h, "Tên danh mục", "Danh mục"), "slug": col(h, "Slug"), "desc": col(h, "Mô tả"), "img": col(h, "Ảnh")}
        for r in rows[hi + 1:]:
            name = (r[ci["name"]] if ci["name"] is not None else None)
            if not name or not str(name).strip() or str(name).strip().isdigit(): continue
            slug = (str(r[ci["slug"]]).strip() if ci["slug"] is not None and r[ci["slug"]] else "") or slugify(name)
            key = norm(name)
            if key in cat_by_key: continue
            new = {"slug": slug, "name": str(name).strip(), "description": str(r[ci["desc"]] or "") if ci["desc"] is not None else "", "image": (str(r[ci["img"]]).strip() if ci["img"] is not None and r[ci["img"]] else None), "count": 0}
            categories.append(new); cat_by_key[key] = slug; cat_by_key[slug] = slug
            print(f"+ category {new['name']} ({slug})")

# --- pricing parameters (optional sheet "Tham số": rate JPY→VND, shipping fee %) used when "Giá vốn" cells are formulas
# that Excel has not recalculated yet (e.g. right after a script wrote "Giá Nhật (JPY)").
RATE, FEE = None, 0.0
if "Tham số" in wb.sheetnames:
    for r in wb["Tham số"].iter_rows(values_only=True):
        cells = [str(x or "") for x in r]
        nums = [x for x in r if isinstance(x, (int, float))]
        label = norm(" ".join(cells))
        if "ty gia" in label and nums: RATE = float(nums[0])
        elif ("phi van chuyen" in label or "van chuyen" in label) and nums: FEE = float(nums[0]); FEE = FEE / 100 if FEE > 1 else FEE
    if RATE: print(f"pricing params: 1 JPY = {RATE} VND, shipping {FEE*100:.0f}%")

# --- products sheet
def find_header(ws):
    rows = list(ws.iter_rows(values_only=True))
    idx = next((i for i, r in enumerate(rows) if any(norm(x).startswith("ten san pham") for x in r if x)), None)
    return rows, idx

if args.sheet:
    rows, hdr_idx = find_header(wb[args.sheet]); sheet = args.sheet
else:
    sheet = None
    for cand in ["Sản phẩm", "Sản phẩm mới"] + wb.sheetnames:
        if cand in wb.sheetnames:
            rows, hdr_idx = find_header(wb[cand])
            if hdr_idx is not None: sheet = cand; break
if sheet is None or hdr_idx is None: sys.exit("no sheet with a 'Tên sản phẩm' header column found")
print(f"reading sheet '{sheet}' (header row {hdr_idx + 1})")
H = [str(x or "") for x in rows[hdr_idx]]
C = {
    "id": next((i for i, x in enumerate(H) if norm(x) == "id"), None), "name": col(H, "Tên sản phẩm"), "slug": col(H, "Đường dẫn"), "cats": col(H, "Danh mục"),
    "price": col(H, "Giá (VNĐ)", "Giá bán (VNĐ)", "Giá bán"), "regular": col(H, "Giá gốc"), "cost": col(H, "Giá vốn"), "sku": col(H, "Mã SKU"),
    "stock": col(H, "Tồn kho"), "oos": col(H, "Hết hàng"), "status": col(H, "Trạng thái"), "tags": col(H, "Từ khóa"),
    "images": col(H, "Danh sách ảnh", "Ảnh (URL", "Ảnh"), "short": col(H, "Mô tả ngắn"), "desc": col(H, "Mô tả chi tiết"),
    "jpy": col(H, "Giá Nhật"), "supplier": col(H, "Link tham khảo", "Link nhà cung cấp", "Link mua"),
    "weight": col(H, "Khối lượng"), "dims": col(H, "Kích thước"), "conf": col(H, "Độ tin cậy", "Tin cậy"), "dsrc": col(H, "Nguồn KT", "Cơ sở", "Nguồn/Cơ sở"),
    "nameJa": col(H, "Tên tiếng Nhật"), "shortJa": col(H, "Mô tả ngắn tiếng Nhật"), "descJa": col(H, "Mô tả tiếng Nhật"),
}
missing = [k for k in ("name",) if C[k] is None]
if missing: sys.exit(f"missing required columns: {missing}; headers seen: {H}")
get = lambda r, k: (r[C[k]] if C[k] is not None and C[k] < len(r) else None)

os.makedirs(IMG_DIR, exist_ok=True)
sess = requests.Session(); sess.headers["User-Agent"] = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/128 Safari/537.36"
now = datetime.datetime.now(datetime.timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z")
next_id = max([2000] + [p["id"] for p in products]) + 1
stats = {"created": 0, "updated": 0, "unchanged": 0, "img_ok": 0, "img_fail": 0}; unknown_cats = set(); warnings = []

def fetch_image(url, fname):
    full = os.path.join(IMG_DIR, fname + ".jpg"); thumb = os.path.join(IMG_DIR, fname + "-300x300.jpg")
    if not os.path.exists(full):
        if args.dry_run: return f"{WEB_DIR}/{fname}.jpg"
        resp = sess.get(url, timeout=30); resp.raise_for_status()
        im = Image.open(io.BytesIO(resp.content)).convert("RGB")
        im.save(full, "JPEG", quality=88, optimize=True)
        ImageOps.pad(im, (300, 300), color=(255, 255, 255)).save(thumb, "JPEG", quality=85, optimize=True)
        time.sleep(0.1)
    return f"{WEB_DIR}/{fname}.jpg"

for r in rows[hdr_idx + 1:]:
    name = get(r, "name")
    if not name or not str(name).strip(): continue
    name = re.sub(r"\s+", " ", str(name)).strip()
    rid = to_int(get(r, "id")); slug_in = (str(get(r, "slug") or "").strip())
    if rid is None and not slug_in:
        name = name.replace("#", "").strip(" .:-–")  # only tidy brand-new rows; existing names are taken as typed
    p = by_id.get(rid) if rid else None
    if p is None and slug_in: p = by_slug.get(slug_in)
    if p is None and not slug_in: p = by_slug.get(slugify(name))   # same name re-imported → same product, not a "-2" copy
    creating = p is None
    if creating:
        slug = slug_in or slugify(name)
        base = slug; n = 2
        while slug in by_slug: slug = f"{base}-{n}"; n += 1
        p = {"id": next_id, "slug": slug, "name": name, "price": 0, "regularPrice": None, "costPrice": None, "currency": "VNĐ", "sku": None,
             "stock": None, "stockStatus": "instock", "categories": [], "tags": [], "images": [], "thumb": "", "shortDescription": "",
             "description": "", "related": [], "rating": None, "reviewCount": 0, "status": "draft", "createdAt": now, "updatedAt": now}
        products.append(p); by_id[p["id"]] = p; by_slug[slug] = p; next_id += 1
    before = json.dumps(p, sort_keys=True, ensure_ascii=False); before_obj = json.loads(before)
    p["name"] = name
    # categories
    cats_raw = get(r, "cats")
    if cats_raw:
        cats = []
        leftovers = []
        for cn in re.split(r"[;\n|]+", str(cats_raw)):
            key = norm(cn)
            if not key: continue
            if key in cat_by_key: cats.append(cat_by_key[key])
            elif cn.strip() in cat_by_key: cats.append(cat_by_key[cn.strip()])
            else: leftovers.append(cn)
        for cn in leftovers:  # legacy comma-separated cell: find every known category name inside it (longest first)
            rest = norm(cn); found = False
            for key in sorted((k for k in cat_by_key if " " in k), key=len, reverse=True):
                if key in rest: cats.append(cat_by_key[key]); rest = rest.replace(key, " "); found = True
            if not found: unknown_cats.add(cn.strip())
        if cats: p["categories"] = list(dict.fromkeys(cats))
    # prices / stock
    price = to_int(get(r, "price"))
    if price is not None: p["price"] = price
    for key, field in (("regular", "regularPrice"), ("cost", "costPrice"), ("stock", "stock")):
        raw = get(r, key)
        if raw is None or str(raw).strip() == "": continue
        p[field] = None if str(raw).strip() == "-" else to_int(raw)
    # cost fallback: Giá vốn empty (formula not yet calculated) but Giá Nhật + rate known → compute here, rounded to 1.000đ
    jpy = to_int(get(r, "jpy"))
    if p.get("costPrice") is None and jpy and RATE:
        p["costPrice"] = int(round(jpy * RATE * (1 + FEE), -3))
    sup = get(r, "supplier")
    if sup and str(sup).strip().startswith("http"): p["supplierUrl"] = str(sup).strip()
    w = get(r, "weight")
    if w not in (None, ""):
        try: p["weightG"] = int(float(str(w).replace(",", ".").replace("g", "").strip()))
        except ValueError: warnings.append(f"row {ri}: khối lượng không hợp lệ: {w}")
    dm = get(r, "dims")
    if dm not in (None, ""):
        p["dimsCm"] = re.sub(r"\s+", "", str(dm)).replace("×", "x").replace("*", "x")
    cf = norm(get(r, "conf"))
    if cf:
        p["dimsConfidence"] = "high" if cf.startswith("cao") or cf == "high" else "medium" if cf.startswith("trung") or cf == "medium" else "low" if cf.startswith("thap") or cf == "low" else p.get("dimsConfidence")
    ds = get(r, "dsrc")
    if ds not in (None, ""): p["dimsSource"] = str(ds).strip()[:300]
    for key, field in (("nameJa", "nameJa"), ("shortJa", "shortDescriptionJa"), ("descJa", "descriptionJa")):
        v = get(r, key)
        if v is not None and str(v).strip(): p[field] = str(v).strip()
    if p.get("regularPrice") is not None and p["regularPrice"] <= p["price"]: p["regularPrice"] = None
    sku = get(r, "sku")
    if sku is not None and str(sku).strip(): p["sku"] = str(sku).strip()
    oos = str(get(r, "oos") or "").strip().lower()
    if oos: p["stockStatus"] = "outofstock" if oos in ("x", "true", "1", "có", "yes") else "instock"
    if p.get("stock") == 0: p["stockStatus"] = "outofstock"
    st = norm(get(r, "status"))
    if st.startswith("dang ban") or st == "publish": p["status"] = "publish"
    elif st.startswith("ban nhap") or st == "draft": p["status"] = "draft"
    if p["price"] <= 0 and p["status"] == "publish":
        p["status"] = "draft"; warnings.append(f"#{p['id']} {name[:40]}: chưa có giá bán → giữ Bản nháp")
    tags = get(r, "tags")
    if tags: p["tags"] = [t.strip() for t in str(tags).split(",") if t.strip()]
    for key, field in (("short", "shortDescription"), ("desc", "description")):
        v = get(r, key)
        if v is not None and str(v).strip(): p[field] = str(v).strip()
    # images
    imgs_raw = get(r, "images")
    if imgs_raw:
        out = []
        for i, u in enumerate([x.strip() for x in re.split(r"[;\n]+", str(imgs_raw)) if x.strip()][:6]):
            if u.startswith("http"):
                # Facebook CDN links expire and were already downloaded on the first import — keep the local copy.
                if "fbcdn.net" in u and any(x.startswith("/sites/") for x in p.get("images", [])):
                    out.extend(x for x in p["images"] if x not in out); continue
                try:
                    out.append(fetch_image(u, f"{p['slug']}{'' if i == 0 else f'-{i+1}'}")); stats["img_ok"] += 1
                except Exception as e:
                    stats["img_fail"] += 1; warnings.append(f"#{p['id']} ảnh lỗi: {str(e)[:60]}")
            else:
                out.append(u)
        if out and out != p.get("images"):
            p["images"] = out
            t = out[0]
            p["thumb"] = t.replace(".jpg", "-300x300.jpg") if t.startswith(WEB_DIR) else t
    if not p.get("thumb") and p.get("images"): p["thumb"] = p["images"][0]
    after = json.dumps(p, sort_keys=True, ensure_ascii=False)
    if creating: stats["created"] += 1
    elif after != before:
        stats["updated"] += 1; p["updatedAt"] = now
        if args.diff:
            changed = [k for k in p if k != "updatedAt" and json.dumps(p[k], ensure_ascii=False) != json.dumps(before_obj.get(k), ensure_ascii=False)]
            print(f"  ~ #{p['id']} {p['name'][:40]}: {', '.join(changed)}")
    else: stats["unchanged"] += 1

if unknown_cats: warnings.append("Danh mục không nhận ra (bỏ qua): " + ", ".join(sorted(unknown_cats)))
seed["meta"]["seededAt"] = now
print(f"created {stats['created']}, updated {stats['updated']}, unchanged {stats['unchanged']}, images ok {stats['img_ok']} / failed {stats['img_fail']}, categories {len(categories)}")
for w in warnings[:40]: print("  !", w)
if args.dry_run:
    print("dry-run: seed.json not written")
else:
    json.dump(seed, open(SEED, "w", encoding="utf-8", newline="\n"), ensure_ascii=False, indent=1)
    print(f"seed.json written ({len(products)} products). Next: npm run dev (seed sync) or commit + tag a release.")
