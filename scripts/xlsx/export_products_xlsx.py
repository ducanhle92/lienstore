"""Export the catalogue (data/seed.json) to an Excel workbook that people (or an AI assistant) can fill in,
then feed back with import_products_xlsx.py.

    python scripts/xlsx/export_products_xlsx.py [output.xlsx]      (default: lienstore-san-pham.xlsx in the repo root)

Sheets:
  "Sản phẩm"  — one row per product, columns = the import contract (see HEADERS below)
  "Danh mục"  — existing categories (name, slug, description, image); add rows to create new categories
  "Hướng dẫn" — column rules for whoever fills the sheet
"""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
SEED = os.path.join(ROOT, "data", "seed.json")
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "lienstore-san-pham.xlsx")

HEADERS = [
    ("ID", 7), ("Tên sản phẩm *", 42), ("Đường dẫn (slug)", 34), ("Danh mục * (tên, cách nhau bằng ;)", 40),
    ("Giá bán (VNĐ) *", 14), ("Giá gốc (VNĐ)", 13), ("Giá vốn (VNĐ)", 13), ("Mã SKU", 12), ("Tồn kho", 9),
    ("Hết hàng (x)", 10), ("Trạng thái", 12), ("Từ khóa (cách nhau bằng ,)", 32),
    ("Ảnh (URL hoặc /sites/..., mỗi ảnh một dòng)", 60), ("Mô tả ngắn (HTML)", 60), ("Mô tả chi tiết (HTML)", 80), ("Ghi chú", 30), ("Link nhà cung cấp", 50),
    ("Khối lượng (g)", 12), ("Kích thước (cm, DxRxC)", 16), ("Độ tin cậy KT (Cao/Trung bình/Thấp)", 14), ("Nguồn KT", 30),
    ("Tên tiếng Nhật", 36), ("Mô tả ngắn tiếng Nhật", 40), ("Mô tả tiếng Nhật (HTML)", 60),
]
CONF = {"high": "Cao", "medium": "Trung bình", "low": "Thấp"}
STATUS = {"publish": "Đang bán", "draft": "Bản nháp"}

seed = json.load(open(SEED, encoding="utf-8"))
cat_name = {c["slug"]: c["name"] for c in seed["categories"]}

wb = Workbook()
ws = wb.active; ws.title = "Sản phẩm"
head_font = Font(bold=True, color="FFFFFF"); head_fill = PatternFill("solid", fgColor="1C7F9E")
need_fill = PatternFill("solid", fgColor="FFF3CD")
for i, (h, w) in enumerate(HEADERS, 1):
    c = ws.cell(row=1, column=i, value=h); c.font = head_font; c.fill = head_fill
    c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[1].height = 32
ws.freeze_panes = "C2"

for r, p in enumerate(sorted(seed["products"], key=lambda x: x["id"]), 2):
    row = [
        p["id"], p["name"], p["slug"], "; ".join(cat_name.get(s, s) for s in p.get("categories", [])),
        p.get("price") or None, p.get("regularPrice"), p.get("costPrice"), p.get("sku"), p.get("stock"),
        "x" if p.get("stockStatus") == "outofstock" else None, STATUS.get(p.get("status"), "Bản nháp"),
        ", ".join(p.get("tags", [])), "\n".join(p.get("images", [])), p.get("shortDescription", ""), p.get("description", ""), None, p.get("supplierUrl"),
        p.get("weightG"), p.get("dimsCm"), CONF.get(p.get("dimsConfidence") or "", None), p.get("dimsSource") or None,
        p.get("nameJa") or None, p.get("shortDescriptionJa") or None, p.get("descriptionJa") or None,
    ]
    for cidx, v in enumerate(row, 1):
        cell = ws.cell(row=r, column=cidx, value=v)
        if cidx in (13, 14, 15): cell.alignment = Alignment(wrap_text=False, vertical="top")
    if not p.get("price"): ws.cell(row=r, column=5).fill = need_fill
    if p.get("costPrice") is None: ws.cell(row=r, column=7).fill = need_fill
dv = DataValidation(type="list", formula1='"Đang bán,Bản nháp"', allow_blank=True); ws.add_data_validation(dv)
dv.add(f"K2:K{ws.max_row + 500}")

wc = wb.create_sheet("Danh mục")
for i, (h, w) in enumerate([("Tên danh mục *", 44), ("Slug", 34), ("Mô tả", 60), ("Ảnh (URL hoặc /sites/...)", 60)], 1):
    c = wc.cell(row=1, column=i, value=h); c.font = head_font; c.fill = head_fill; wc.column_dimensions[get_column_letter(i)].width = w
for r, c in enumerate(seed["categories"], 2):
    wc.append([c["name"], c["slug"], c.get("description", ""), c.get("image")])

wg = wb.create_sheet("Hướng dẫn")
wg.column_dimensions["A"].width = 120
for line in [
    "QUY TẮC ĐIỀN SHEET 'Sản phẩm' (script scripts/xlsx/import_products_xlsx.py đọc theo tên cột, không theo vị trí)",
    "• ID: giữ nguyên cho sản phẩm có sẵn. Dòng mới để trống ID (script tự cấp).",
    "• Tên sản phẩm *: bắt buộc. Đường dẫn (slug): để trống sẽ tạo từ tên; KHÔNG đổi slug của sản phẩm cũ (mất link).",
    "• Danh mục *: tên đúng như sheet 'Danh mục', nhiều danh mục cách nhau bằng dấu ';'. Tên chưa có → thêm dòng vào sheet 'Danh mục' trước.",
    "• Giá bán / Giá gốc / Giá vốn: số nguyên VNĐ, không dấu chấm, không chữ (vd 350000). Ô vàng = còn thiếu.",
    "• Giá gốc chỉ điền khi đang giảm giá và phải lớn hơn Giá bán. Giá vốn chỉ hiển thị trong trang quản trị (tính lợi nhuận).",
    "• Tồn kho: số nguyên; để trống = không theo dõi. Hết hàng: đánh 'x'.",
    "• Trạng thái: 'Đang bán' hoặc 'Bản nháp'. Sản phẩm chưa có Giá bán sẽ bị ép về 'Bản nháp'.",
    "• Từ khóa: cách nhau bằng dấu ','. Ảnh: mỗi ảnh một dòng (Alt+Enter) hoặc cách nhau bằng ';'; URL http(s) sẽ được tải về kho ảnh web, ảnh đầu là ảnh đại diện.",
    "• Mô tả ngắn / chi tiết: HTML đơn giản (<p>, <ul><li>, <strong>, <br>). Không dán script/iframe.",
    "• Không xoá dòng để xoá sản phẩm (script không xoá). Muốn ẩn → đặt 'Bản nháp'.",
]:
    wg.append([line])
wg["A1"].font = Font(bold=True)

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print(f"exported {len(seed['products'])} products, {len(seed['categories'])} categories → {OUT}")
